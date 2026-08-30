import os
import sys

# Ensure the project root is in sys.path so `from src.xxx` imports work
# regardless of whether this file is run directly or as a module.
_project_root = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if _project_root not in sys.path:
    sys.path.insert(0, _project_root)

import tkinter as tk
from tkinter import filedialog, messagebox
import threading
import customtkinter as ctk

from src.config import *
from src.operations import (
    op_truncate_after_n, op_rename_ext, op_convert_to_jpg, op_pdf_to_jpg, op_rotate_images,
    op_rename_sequence, op_extract_photos, op_extract_photos_dlib, op_extract_photos_hybrid,
    op_excel_rename, build_excel_rename_plan
)
from src.file_explorer import FileExplorerGrid

class TruncatePreviewModal(ctk.CTkToplevel):
    """Karakter kesme öncesi değişiklikleri gösteren önizleme penceresi."""
    def __init__(self, master, file_paths, n, **kwargs):
        super().__init__(master, **kwargs)
        self.title("Kesme Önizlemesi")
        self.geometry("600x450")
        
        # Dialogu üstte ve modal tut
        # Windows bazen topmost'a rağmen arkaya atabilir, transient o yüzden önemli.
        self.transient(master)
        self.grab_set() 
        
        self.grid_columnconfigure(0, weight=1)
        self.grid_rowconfigure(1, weight=1)
        
        self.result = False
        
        # Başlık
        ctk.CTkLabel(
            self, text=f"İlk {n} karakter sonrası silinecek. Onaylıyor musunuz?", 
            font=FONT_HEAD, text_color=TEXT
        ).grid(row=0, column=0, pady=15, padx=20, sticky="w")
        
        # Kaydırılabilir liste
        scroll = ctk.CTkScrollableFrame(self, fg_color=CARD, corner_radius=10)
        scroll.grid(row=1, column=0, sticky="nsew", padx=20, pady=(0, 10))
        
        for path in file_paths:
            if not os.path.isfile(path): continue
            
            fname = os.path.basename(path)
            name, ext = os.path.splitext(fname)
            
            if len(name) > n:
                new_name = name[:n] + ext
                lbl_text = f"📄 {fname}   ➔   {new_name}"
                color = TEXT
            else:
                lbl_text = f"📄 {fname}   ➔   (Değişmeyecek)"
                color = SUBTEXT
                
            ctk.CTkLabel(scroll, text=lbl_text, font=FONT_MONO, text_color=color).pack(anchor="w", pady=2, padx=5)
            
        # Alt butonlar
        btn_frame = ctk.CTkFrame(self, fg_color="transparent")
        btn_frame.grid(row=2, column=0, pady=15, padx=20, sticky="e")
        
        ctk.CTkButton(
            btn_frame, text="İptal", fg_color=BORDER, hover_color=CARD, 
            text_color=TEXT, font=FONT_BODY, command=self._cancel, width=100
        ).pack(side="left", padx=10)
        
        ctk.CTkButton(
            btn_frame, text="Onayla ve Çalıştır", fg_color=ACCENT, hover_color=ACCENT2, 
            text_color="#ffffff", font=FONT_BODY, command=self._confirm, width=150
        ).pack(side="left")
        
    def _cancel(self):
        self.result = False
        self.destroy()
        
    def _confirm(self):
        self.result = True
        self.destroy()

class ExcelRenameModal(ctk.CTkToplevel):
    """Excel ile imaj numaralandırma ayarlarını toplayan modal pencere."""
    def __init__(self, master, file_paths=None, **kwargs):
        super().__init__(master, **kwargs)
        self.title("📊 Excel ile Numaralandır")
        self.geometry("760x650")
        self.transient(master)
        self.grab_set()
        self.protocol("WM_DELETE_WINDOW", self._cancel)

        self.grid_columnconfigure(0, weight=1)
        self.grid_rowconfigure(5, weight=1)  # Önizleme alanı genişlesin

        self.result = False
        self.excel_path = ""
        self.sheet_name = ""
        self.column_letter = ""
        self.start_num = 1
        self.output_column = ""
        self.mode = "match"
        self._file_paths = file_paths or []

        self._excel_var = tk.StringVar(value="")
        self._sheet_var = tk.StringVar(value="")
        self._col_var = tk.StringVar(value="")
        self._start_var = tk.StringVar(value="100")
        self._output_col_var = tk.StringVar(value="")
        self._folder_mode_var = tk.BooleanVar(value=False)
        self._wb = None  # openpyxl workbook referansı

        # ── Başlık ──
        ctk.CTkLabel(
            self, text="📊  İmaj Dosyalarını Excel ile Numaralandır",
            font=FONT_HEAD, text_color=TEXT
        ).grid(row=0, column=0, pady=(15, 10), padx=20, sticky="w")

        # ── Excel Dosya Seçimi ──
        excel_frame = ctk.CTkFrame(self, fg_color=CARD, corner_radius=8)
        excel_frame.grid(row=1, column=0, sticky="ew", padx=20, pady=(0, 8))
        excel_frame.columnconfigure(1, weight=1)

        ctk.CTkLabel(excel_frame, text="Excel Dosyası:", font=FONT_BODY, text_color=SUBTEXT).grid(
            row=0, column=0, padx=(10, 6), pady=10, sticky="w")
        self._excel_entry = ctk.CTkEntry(
            excel_frame, textvariable=self._excel_var, font=FONT_MONO, height=30,
            fg_color=BG, border_color=BORDER, text_color=TEXT,
            placeholder_text="Excel dosyası seçin..."
        )
        self._excel_entry.grid(row=0, column=1, padx=(0, 6), pady=10, sticky="ew")
        ctk.CTkButton(
            excel_frame, text="Seç", width=60, height=30, font=FONT_BODY,
            fg_color=ACCENT, hover_color=ACCENT2, command=self._pick_excel
        ).grid(row=0, column=2, padx=(0, 10), pady=10)

        # ── Sekme ve Sütun Seçimi ──
        sel_frame = ctk.CTkFrame(self, fg_color=CARD, corner_radius=8)
        sel_frame.grid(row=2, column=0, sticky="ew", padx=20, pady=(0, 8))
        sel_frame.columnconfigure(1, weight=1)
        sel_frame.columnconfigure(3, weight=1)

        ctk.CTkLabel(sel_frame, text="Sekme:", font=FONT_BODY, text_color=SUBTEXT).grid(
            row=0, column=0, padx=(10, 6), pady=10, sticky="w")
        self._sheet_menu = ctk.CTkOptionMenu(
            sel_frame, variable=self._sheet_var, values=["—"],
            font=FONT_BODY, fg_color=BG, button_color=ACCENT,
            button_hover_color=ACCENT2, text_color=TEXT,
            command=self._on_sheet_changed
        )
        self._sheet_menu.grid(row=0, column=1, padx=(0, 12), pady=10, sticky="ew")

        ctk.CTkLabel(sel_frame, text="Adların Bulunduğu Sütun:", font=FONT_BODY, text_color=SUBTEXT).grid(
            row=0, column=2, padx=(10, 6), pady=10, sticky="w")
        self._col_menu = ctk.CTkOptionMenu(
            sel_frame, variable=self._col_var, values=["—"],
            font=FONT_BODY, fg_color=BG, button_color=ACCENT,
            button_hover_color=ACCENT2, text_color=TEXT,
            command=self._on_col_changed
        )
        self._col_menu.grid(row=0, column=3, padx=(0, 10), pady=10, sticky="ew")

        ctk.CTkLabel(sel_frame, text="Yazılacak Sütun:", font=FONT_BODY, text_color=SUBTEXT).grid(
            row=1, column=0, padx=(10, 6), pady=(0, 10), sticky="w")
        self._output_col_menu = ctk.CTkOptionMenu(
            sel_frame, variable=self._output_col_var, values=["—"],
            font=FONT_BODY, fg_color=BG, button_color=ACCENT,
            button_hover_color=ACCENT2, text_color=TEXT,
            command=lambda _value: self._update_preview()
        )
        self._output_col_menu.grid(row=1, column=1, columnspan=3, padx=(0, 10), pady=(0, 10), sticky="ew")

        # ── Başlangıç Numarası ──
        num_frame = ctk.CTkFrame(self, fg_color=CARD, corner_radius=8)
        num_frame.grid(row=3, column=0, sticky="ew", padx=20, pady=(0, 8))

        ctk.CTkLabel(num_frame, text="Başlangıç Numarası:", font=FONT_BODY, text_color=SUBTEXT).pack(
            side="left", padx=(10, 6), pady=10)
        self._start_entry = ctk.CTkEntry(
            num_frame, textvariable=self._start_var, width=80, height=30,
            font=FONT_BODY, fg_color=BG, border_color=BORDER, text_color=TEXT
        )
        self._start_entry.pack(side="left", padx=(0, 10), pady=10)
        self._start_entry.bind("<KeyRelease>", lambda _event: self._update_preview())
        ctk.CTkCheckBox(
            num_frame, text="Excel ile eşleştirmeden klasör dosyalarını sırayla adlandır",
            variable=self._folder_mode_var, command=self._update_preview,
            font=FONT_BODY, fg_color=ACCENT, hover_color=ACCENT2
        ).pack(side="left", padx=10, pady=10)

        # ── Önizleme Başlığı ──
        ctk.CTkLabel(
            self, text="Önizleme (en az %70 Türkçe ad benzerliği):",
            font=FONT_BODY, text_color=SUBTEXT
        ).grid(row=4, column=0, pady=(4, 2), padx=20, sticky="w")

        # ── Önizleme Alanı ──
        self._preview_scroll = ctk.CTkScrollableFrame(self, fg_color=CARD, corner_radius=8)
        self._preview_scroll.grid(row=5, column=0, sticky="nsew", padx=20, pady=(0, 10))

        # ── Alt Butonlar ──
        btn_frame = ctk.CTkFrame(self, fg_color="transparent")
        btn_frame.grid(row=6, column=0, pady=12, padx=20, sticky="e")

        ctk.CTkButton(
            btn_frame, text="İptal", fg_color=BORDER, hover_color=CARD,
            text_color=TEXT, font=FONT_BODY, command=self._cancel, width=100
        ).pack(side="left", padx=10)

        ctk.CTkButton(
            btn_frame, text="Onayla ve Çalıştır", fg_color=ACCENT, hover_color=ACCENT2,
            text_color="#ffffff", font=FONT_BODY, command=self._confirm, width=150
        ).pack(side="left")

    def _pick_excel(self):
        path = filedialog.askopenfilename(
            title="Excel Dosyası Seç",
            # openpyxl eski .xls biçimini desteklemez; seçilebilir göstermek
            # kullanıcıya çalışacakmış izlenimi veriyordu.
            filetypes=[("Excel Dosyaları", "*.xlsx"), ("Tüm Dosyalar", "*.*")]
        )
        if not path:
            return
        self._excel_var.set(path)
        try:
            from openpyxl import load_workbook
            self._wb = load_workbook(path, read_only=False)
            sheets = self._wb.sheetnames
            self._sheet_menu.configure(values=sheets)
            if sheets:
                self._sheet_var.set(sheets[0])
                self._on_sheet_changed(sheets[0])
        except Exception as e:
            messagebox.showerror("Hata", f"Excel dosyası açılamadı:\n{e}", parent=self)

    def _on_sheet_changed(self, sheet_name):
        if self._wb is None:
            return
        try:
            from openpyxl.utils import get_column_letter
            ws = self._wb[sheet_name]
            cols = []
            for col in range(1, ws.max_column + 1):
                letter = get_column_letter(col)
                header = ws.cell(row=1, column=col).value
                if header:
                    cols.append(f"{letter} ({header})")
                else:
                    cols.append(letter)
            if cols:
                self._col_menu.configure(values=cols)
                self._col_var.set(cols[0])
                self._output_col_menu.configure(values=cols)
                output = cols[1] if len(cols) > 1 else cols[0]
                self._output_col_var.set(output)
                self._on_col_changed(cols[0])
            else:
                self._col_menu.configure(values=["—"])
                self._col_var.set("—")
        except Exception:
            pass

    def _on_col_changed(self, col_display):
        """Sütun değiştiğinde önizlemeyi güncelle."""
        self._update_preview()

    def _update_preview(self):
        """Mevcut ayarlarla önizleme listesini günceller."""
        # Önceki önizlemeyi temizle
        for widget in self._preview_scroll.winfo_children():
            widget.destroy()

        if self._wb is None:
            return

        sheet_name = self._sheet_var.get()
        col_display = self._col_var.get()
        if not sheet_name or col_display == "—":
            return

        # Sütun harfini çıkar ("A (Başlık)" -> "A")
        col_letter = col_display.split(" ")[0].strip()

        try:
            start_num = int(self._start_var.get().strip())
        except ValueError:
            ctk.CTkLabel(self._preview_scroll, text="⚠ Geçersiz başlangıç numarası",
                        font=FONT_MONO, text_color=DANGER).pack(anchor="w", pady=2, padx=5)
            return

        try:
            from openpyxl.utils import column_index_from_string
            mode = "folder" if self._folder_mode_var.get() else "match"
            plan = build_excel_rename_plan(self._file_paths, self._excel_var.get(),
                                           sheet_name, col_letter, start_num, mode)
            preview_count = 0
            for item in plan[:50]:
                old_name = os.path.basename(item["file_path"])
                ext = os.path.splitext(old_name)[1]
                similarity_text = (f" — %{item['similarity']:.1f} benzer"
                                   if "similarity" in item else "")
                preview_color = (DANGER if item.get("similarity", 100) < 80 else TEXT)
                lbl_text = (f"📄 {old_name}  ➔  {item['new_num']}{ext}   "
                            f"(Excel satır {item['row']}, ad: {item['cell_value']}{similarity_text})")
                ctk.CTkLabel(self._preview_scroll, text=lbl_text,
                            font=FONT_MONO, text_color=preview_color).pack(anchor="w", pady=2, padx=5)
                preview_count += 1

            if preview_count == 0:
                ctk.CTkLabel(self._preview_scroll, text="⚠ Dosya ve Excel verileri arasında eşleşme bulunamadı.",
                            font=FONT_MONO, text_color=WARNING).pack(anchor="w", pady=2, padx=5)
        except Exception as e:
            ctk.CTkLabel(self._preview_scroll, text=f"✖ Önizleme hatası: {e}",
                        font=FONT_MONO, text_color=DANGER).pack(anchor="w", pady=2, padx=5)

    def _cancel(self):
        self.result = False
        if self._wb:
            self._wb.close()
            self._wb = None
        self.destroy()

    def _confirm(self):
        # Doğrulamalar
        if not self._excel_var.get().strip():
            messagebox.showwarning("Uyarı", "Lütfen bir Excel dosyası seçin.", parent=self)
            return
        if self._sheet_var.get() == "—" or not self._sheet_var.get():
            messagebox.showwarning("Uyarı", "Lütfen bir sekme seçin.", parent=self)
            return
        if self._col_var.get() == "—" or not self._col_var.get():
            messagebox.showwarning("Uyarı", "Lütfen bir sütun seçin.", parent=self)
            return
        if self._output_col_var.get() == "—" or not self._output_col_var.get():
            messagebox.showwarning("Uyarı", "Lütfen yeni adların yazılacağı sütunu seçin.", parent=self)
            return
        source_column = self._col_var.get().split(" ")[0].strip()
        output_column = self._output_col_var.get().split(" ")[0].strip()
        if source_column == output_column:
            messagebox.showwarning(
                "Sütun Seçimi",
                "Adların korunması için kaynak sütun ile numaranın yazılacağı sütun farklı olmalıdır.",
                parent=self
            )
            return
        try:
            start = int(self._start_var.get().strip())
            if start < 0:
                raise ValueError
        except ValueError:
            messagebox.showerror("Hata", "Lütfen geçerli bir başlangıç numarası girin.", parent=self)
            return

        mode = "folder" if self._folder_mode_var.get() else "match"
        plan = build_excel_rename_plan(
            self._file_paths, self._excel_var.get().strip(), self._sheet_var.get(),
            self._col_var.get().split(" ")[0].strip(), start, mode
        )
        if not plan:
            messagebox.showwarning("Eşleşme Yok", "İşlem yapılabilecek dosya bulunamadı.", parent=self)
            return
        if not messagebox.askyesno(
            "Eşleştirmeyi Onayla",
            f"En az %70 benzerliğe sahip {len(plan)} eşleşme bulundu.\n"
            f"Dosyalar numarayla adlandırılacak ve numaralar {output_column} sütununa yazılacak.\n\n"
            "Önizlemedeki eşleştirmeleri onaylıyor musunuz?",
            parent=self
        ):
            return

        self.result = True
        self.excel_path = self._excel_var.get().strip()
        self.sheet_name = self._sheet_var.get()
        self.column_letter = self._col_var.get().split(" ")[0].strip()
        self.start_num = start
        self.output_column = output_column
        self.mode = mode

        if self._wb:
            self._wb.close()
            self._wb = None
        self.destroy()


class AllToolsApp(ctk.CTk):
    def __init__(self):
        super().__init__()
        init_theme()
        self.title("dESkart dosya düzenleme v1.7")
        self.geometry("1100x800")
        self.minsize(800, 600)
        self.configure(fg_color=BG)
        
        # Pencere ikonu ayarı
        try:
            if os.path.exists("fav.ico"):
                self.iconbitmap("fav.ico")
        except Exception:
            pass

        self._folder = tk.StringVar(value="")
        self._n_var  = tk.StringVar(value="4")
        self._dpi_var = tk.StringVar(value="150")
        self._apply_to_all = tk.BooleanVar(value=False)
        self._hybrid_start_var = tk.StringVar(value="1")

        # İşlem seçenek değişkenleri
        self._opts = {
            "truncate":   tk.BooleanVar(),
            "jpeg2jpg":   tk.BooleanVar(),
            "bmp2jpg":    tk.BooleanVar(),
            "png2jpg":    tk.BooleanVar(),
            "jfif2jpg":   tk.BooleanVar(),
            "pdf2jpg":    tk.BooleanVar(),
            "ext_photos": tk.BooleanVar(),
            "ext_photos_dlib": tk.BooleanVar(),
            "ext_photos_hybrid": tk.BooleanVar(),
            "seq_rename": tk.BooleanVar(),
            "excel_rename": tk.BooleanVar(),
        }
        self._excel_rename_params = None

        # Ardışık isim değiştirme başlangıç numarası
        self._seq_start_var = tk.StringVar(value="1")

        self._build_ui()
        self.after(500, self._init_windnd)

    def _init_windnd(self):
        try:
            import windnd
            # İşletim sistemi dnd olayını ana pencereye bağlıyoruz
            windnd.hook_dropfiles(self.winfo_id(), func=lambda files: self._file_explorer._on_drag_drop(files))
        except Exception as e:
            print(f"windnd hook failed: {e}")

    # ── UI İnşası ──────────────────────────────────
    def _build_ui(self):
        # ── Üst başlık ──
        header = ctk.CTkFrame(self, fg_color=PANEL, corner_radius=0, height=64)
        header.pack(fill="x", side="top")
        header.pack_propagate(False)

        ctk.CTkLabel(
            header, text="⚙  dESkart",
            font=FONT_TITLE, text_color=TEXT
        ).pack(side="left", padx=24, pady=14)

        ctk.CTkLabel(
            header, text="dosya düzenleme v1.7",
            font=FONT_BODY, text_color=SUBTEXT
        ).pack(side="left", padx=0, pady=14)

        # ── Ana içerik ──
        body = ctk.CTkFrame(self, fg_color=BG)
        body.pack(fill="both", expand=True, padx=20, pady=(16, 20))
        body.columnconfigure(0, weight=0, minsize=320) # Sol panel sabit minimum genişlikte
        body.columnconfigure(1, weight=1)
        body.rowconfigure(0, weight=1)

        # Sol panel — tüm form işlemleri
        left = ctk.CTkScrollableFrame(body, fg_color=PANEL, corner_radius=14)
        left.grid(row=0, column=0, sticky="nsew", padx=(0,10), pady=0)
        self._build_left_panel(left)

        # Sağ panel — sekmeler (Dosyalar ve Log)
        self._tab_view = ctk.CTkTabview(body, fg_color=PANEL, corner_radius=14)
        self._tab_view.grid(row=0, column=1, sticky="nsew", pady=(0, 0))
        self._tab_view.add("Dosya Gezgini")
        self._tab_view.add("İşlem Günlüğü")

        self._build_file_explorer(self._tab_view.tab("Dosya Gezgini"))
        self._build_log(self._tab_view.tab("İşlem Günlüğü"))

    def _build_left_panel(self, parent):
        parent.columnconfigure(0, weight=1)

        # -- Klasör Seçimi --
        folder_frame = ctk.CTkFrame(parent, fg_color="transparent")
        folder_frame.pack(fill="x", padx=12, pady=(15, 8))
        
        ctk.CTkLabel(
            folder_frame, text="📁  Klasör Seçimi",
            font=FONT_HEAD, text_color=TEXT
        ).pack(anchor="w", pady=(0, 4))
        
        entry = ctk.CTkEntry(
            folder_frame, textvariable=self._folder,
            font=FONT_MONO, height=34,
            fg_color=CARD, border_color=BORDER,
            text_color=TEXT, placeholder_text="Açmak için klasör seçin..."
        )
        entry.pack(side="left", fill="x", expand=True, padx=(0, 8))

        ctk.CTkButton(
            folder_frame, text="Seç", width=60, height=34,
            font=FONT_BODY, fg_color=ACCENT, hover_color=ACCENT2,
            command=self._pick_folder
        ).pack(side="right")

        ctk.CTkFrame(parent, fg_color=BORDER, height=1).pack(fill="x", padx=12, pady=(4, 8))

        # -- Uygula Alanı --
        apply_opt_frame = ctk.CTkFrame(parent, fg_color="transparent")
        apply_opt_frame.pack(fill="x", padx=12, pady=(0, 8))
        
        ctk.CTkSwitch(
            apply_opt_frame, text="Tüm Dosyalara Uygula",
            variable=self._apply_to_all, font=FONT_BODY, text_color=TEXT,
            progress_color=ACCENT, button_color=TEXT, button_hover_color=SUBTEXT
        ).pack(side="left")

        self._status = ctk.CTkLabel(
            parent, text="Hazır", font=FONT_BODY, text_color=SUBTEXT
        )
        self._status.pack(pady=(0, 4))
        
        ctk.CTkButton(
            parent,
            text="▶  Seçili Dosyalara Uygula",
            font=("Segoe UI", 13, "bold"),
            height=40,
            fg_color=ACCENT, hover_color=ACCENT2,
            corner_radius=8,
            command=self._run
        ).pack(fill="x", padx=12, pady=(0, 8))

        ctk.CTkFrame(parent, fg_color=BORDER, height=1).pack(fill="x", padx=12, pady=(4, 8))

        # -- İşlemler --
        ctk.CTkLabel(
            parent, text="İşlemler (Seçili Dosyalara)",
            font=FONT_HEAD, text_color=TEXT
        ).pack(anchor="w", padx=12, pady=(0, 4))

        # 1. DOSYA İSİM İŞLEMLERİ (Karakter Kes & Numara Ver)
        name_frame = ctk.CTkFrame(parent, fg_color=CARD, corner_radius=8)
        name_frame.pack(fill="x", padx=12, pady=4)
        
        t_row = ctk.CTkFrame(name_frame, fg_color="transparent")
        t_row.pack(fill="x", padx=10, pady=(10, 4))
        ctk.CTkCheckBox(
            t_row, text="N. karakterden sonrasını sil N =", variable=self._opts["truncate"],
            checkbox_width=20, checkbox_height=20, font=FONT_BODY, fg_color=ACCENT, hover_color=ACCENT2
        ).pack(side="left")
        ctk.CTkEntry(
            t_row, textvariable=self._n_var,
            width=50, height=26, font=FONT_BODY,
            fg_color=BG, border_color=BORDER, text_color=TEXT
        ).pack(side="left", padx=6)

        s_row = ctk.CTkFrame(name_frame, fg_color="transparent")
        s_row.pack(fill="x", padx=10, pady=(4, 10))
        ctk.CTkCheckBox(
            s_row, text="Numaralandırarak Adlandır. Başlangıç =", variable=self._opts["seq_rename"],
            checkbox_width=20, checkbox_height=20, font=FONT_BODY, fg_color=ACCENT, hover_color=ACCENT2
        ).pack(side="left")
        ctk.CTkEntry(
            s_row, textvariable=self._seq_start_var,
            width=50, height=26, font=FONT_BODY,
            fg_color=BG, border_color=BORDER, text_color=TEXT
        ).pack(side="left", padx=6)

        # 2. FORMAT DÖNÜŞTÜRME
        conv_frame = ctk.CTkFrame(parent, fg_color=CARD, corner_radius=8)
        conv_frame.pack(fill="x", padx=12, pady=4)
        
        row1 = ctk.CTkFrame(conv_frame, fg_color="transparent")
        row1.pack(fill="x", padx=10, pady=(10, 4))
        ctk.CTkCheckBox(
            row1, text="🖼 .jpeg → .jpg", variable=self._opts["jpeg2jpg"], font=FONT_BODY,
            checkbox_width=20, checkbox_height=20, fg_color=ACCENT, hover_color=ACCENT2
        ).pack(side="left", expand=True, anchor="w")
        ctk.CTkCheckBox(
            row1, text="🖼 .png → .jpg", variable=self._opts["png2jpg"], font=FONT_BODY,
            checkbox_width=20, checkbox_height=20, fg_color=ACCENT, hover_color=ACCENT2
        ).pack(side="left", expand=True, anchor="w")
        
        row2 = ctk.CTkFrame(conv_frame, fg_color="transparent")
        row2.pack(fill="x", padx=10, pady=4)
        ctk.CTkCheckBox(
            row2, text="🖼 .bmp → .jpg", variable=self._opts["bmp2jpg"], font=FONT_BODY,
            checkbox_width=20, checkbox_height=20, fg_color=ACCENT, hover_color=ACCENT2
        ).pack(side="left", expand=True, anchor="w")
        ctk.CTkCheckBox(
            row2, text="🖼 .jfif → .jpg", variable=self._opts["jfif2jpg"], font=FONT_BODY,
            checkbox_width=20, checkbox_height=20, fg_color=ACCENT, hover_color=ACCENT2
        ).pack(side="left", expand=True, anchor="w")

        pdf_row = ctk.CTkFrame(conv_frame, fg_color="transparent")
        pdf_row.pack(fill="x", padx=10, pady=(4, 10))
        ctk.CTkCheckBox(
            pdf_row, text="📄 .pdf → .jpg", variable=self._opts["pdf2jpg"], font=FONT_BODY,
            checkbox_width=20, checkbox_height=20, fg_color=ACCENT, hover_color=ACCENT2
        ).pack(side="left")
        ctk.CTkLabel(pdf_row, text="DPI:", font=FONT_SMALL, text_color=SUBTEXT).pack(side="left", padx=(10, 4))
        ctk.CTkEntry(
            pdf_row, textvariable=self._dpi_var,
            width=50, height=26, font=FONT_BODY,
            fg_color=BG, border_color=BORDER, text_color=TEXT
        ).pack(side="left", padx=0)

        # 3. EXCEL İLE NUMARALANDIR
        excel_frame2 = ctk.CTkFrame(parent, fg_color=CARD, corner_radius=8)
        excel_frame2.pack(fill="x", padx=12, pady=4)

        e_row = ctk.CTkFrame(excel_frame2, fg_color="transparent")
        e_row.pack(fill="x", padx=10, pady=10)
        ctk.CTkCheckBox(
            e_row, text="📊 Excel ile Numaralandır", variable=self._opts["excel_rename"],
            checkbox_width=20, checkbox_height=20, font=FONT_BODY, fg_color=ACCENT,
            hover_color=ACCENT2, command=self._on_excel_rename_toggled
        ).pack(side="left")
        ctk.CTkLabel(
            e_row, text="(İmaj → Excel eşleştirme)",
            font=FONT_SMALL, text_color=SUBTEXT
        ).pack(side="left", padx=(8, 0))

        # 4. YÜZ KESİCİ SEÇENEKLERİ
        face_frame = ctk.CTkFrame(parent, fg_color=CARD, corner_radius=8)
        face_frame.pack(fill="x", padx=12, pady=4)
        
        ctk.CTkLabel(face_frame, text="✂️ Vesikalıkları Kes/Ayır:", font=FONT_BODY, text_color=SUBTEXT).pack(anchor="w", padx=10, pady=(8,4))
        
        f_row = ctk.CTkFrame(face_frame, fg_color="transparent")
        f_row.pack(fill="x", padx=10, pady=(0, 4))
        ctk.CTkCheckBox(
            f_row, text="Haar Cascade", variable=self._opts["ext_photos"],
            checkbox_width=20, checkbox_height=20, font=FONT_BODY, fg_color=ACCENT, hover_color=ACCENT2
        ).pack(side="left", expand=True, anchor="w")
        ctk.CTkCheckBox(
            f_row, text="Dlib", variable=self._opts["ext_photos_dlib"],
            checkbox_width=20, checkbox_height=20, font=FONT_BODY, fg_color=ACCENT, hover_color=ACCENT2
        ).pack(side="left", expand=True, anchor="w")
        
        f_row2 = ctk.CTkFrame(face_frame, fg_color="transparent")
        f_row2.pack(fill="x", padx=10, pady=(0, 10))
        ctk.CTkCheckBox(
            f_row2, text="Hibrit (Haar + Dlib)", variable=self._opts["ext_photos_hybrid"],
            checkbox_width=20, checkbox_height=20, font=FONT_BODY, fg_color=ACCENT, hover_color=ACCENT2
        ).pack(side="left", anchor="w")
        ctk.CTkLabel(f_row2, text="Başlat:", font=FONT_SMALL, text_color=SUBTEXT).pack(side="left", padx=(8, 2))
        ctk.CTkEntry(
            f_row2, textvariable=self._hybrid_start_var,
            width=46, height=24, font=FONT_BODY,
            fg_color=BG, border_color=BORDER, text_color=TEXT
        ).pack(side="left")

        # Tümünü seç / kaldır
        btn_row = ctk.CTkFrame(parent, fg_color="transparent")
        btn_row.pack(fill="x", padx=12, pady=(10,12))
        ctk.CTkButton(
            btn_row, text="Tümünü Seç", height=30, font=FONT_SMALL,
            fg_color=CARD, hover_color=BORDER, text_color=SUBTEXT,
            command=self._select_all
        ).pack(side="left", expand=True, fill="x", padx=(0,4))
        ctk.CTkButton(
            btn_row, text="Temizle", height=30, font=FONT_SMALL,
            fg_color=CARD, hover_color=BORDER, text_color=SUBTEXT,
            command=self._deselect_all
        ).pack(side="left", expand=True, fill="x", padx=(4,0))

    def _build_log(self, parent):
        parent.rowconfigure(1, weight=1)
        parent.columnconfigure(0, weight=1)

        hdr = ctk.CTkFrame(parent, fg_color="transparent")
        hdr.grid(row=0, column=0, sticky="ew", padx=14, pady=(12,4))
        ctk.CTkLabel(hdr, text="📋  İşlem Günlüğü", font=FONT_HEAD, text_color=TEXT).pack(side="left")
        ctk.CTkButton(
            hdr, text="Log Temizle", width=70, height=26,
            font=FONT_SMALL, fg_color=CARD, hover_color=BORDER, text_color=SUBTEXT,
            command=self._clear_log
        ).pack(side="right")

        self._log_box = ctk.CTkTextbox(
            parent, font=FONT_MONO, fg_color=BG,
            text_color=TEXT, border_color=BORDER,
            scrollbar_button_color=BORDER,
            activate_scrollbars=True
        )
        self._log_box.grid(row=1, column=0, sticky="nsew", padx=12, pady=(0,12))
        self._log_box.configure(state="disabled")

    def _build_file_explorer(self, parent):
        parent.rowconfigure(1, weight=1)
        parent.columnconfigure(0, weight=1)

        # --- Üst Bilgi Barı ---
        hdr = ctk.CTkFrame(parent, fg_color="transparent")
        hdr.grid(row=0, column=0, sticky="ew", padx=14, pady=(12,4))
        
        info_label = ctk.CTkLabel(
            hdr, 
            text="Dosya veya klasörlere sağ tıklayarak işlem yapın (Çoklu seçim: Ctrl/Shift + Tıkla)", 
            font=FONT_SMALL, text_color=SUBTEXT
        )
        info_label.pack(side="left")
        
        ctk.CTkButton(
            hdr, text="Yenile", width=70, height=26,
            font=FONT_SMALL, fg_color=CARD, hover_color=BORDER, text_color=TEXT,
            command=self._refresh_file_list
        ).pack(side="right")

        # --- Grid Tabanlı Dosya Gezgini ---
        self._file_explorer = FileExplorerGrid(parent)
        self._file_explorer.grid(row=1, column=0, sticky="nsew", padx=12, pady=(0,12))

    # ── Yardımcı ──────────────────────────────────
    def _pick_folder(self):
        path = filedialog.askdirectory(title="Klasör Seç")
        if path:
            self._folder.set(path)
            self._log(f"📁 Klasör Açıldı: {path}\n")
            self._refresh_file_list()
            self._tab_view.set("Dosya Gezgini")

    def _refresh_file_list(self):
        folder = self._folder.get().strip()
        if not folder or not os.path.isdir(folder):
            return
        
        # Seçilen yeni klasörü veya mevcut klasörü yenilerken önbelleği temizle
        from src.file_explorer import FileItemCache
        FileItemCache._cache.clear()
        
        self._file_explorer.populate(folder)

    def _select_all(self):
        for v in self._opts.values():
            v.set(True)

    def _deselect_all(self):
        for v in self._opts.values():
            v.set(False)
        self._excel_rename_params = None

    def _log(self, msg: str):
        self._log_box.configure(state="normal")
        self._log_box.insert("end", msg + "\n")
        self._log_box.see("end")
        self._log_box.configure(state="disabled")

    def _clear_log(self):
        self._log_box.configure(state="normal")
        self._log_box.delete("1.0", "end")
        self._log_box.configure(state="disabled")

    def _set_status(self, msg, color=SUBTEXT):
        self._status.configure(text=msg, text_color=color)

    def _on_excel_rename_toggled(self):
        """Excel seçeneği açıldığında gerekli ayarları hemen topla."""
        if not self._opts["excel_rename"].get():
            self._excel_rename_params = None
            return

        folder = self._folder.get().strip()
        if self._apply_to_all.get() and os.path.isdir(folder):
            preview_files = [os.path.join(folder, name) for name in os.listdir(folder)
                             if os.path.isfile(os.path.join(folder, name))]
        else:
            preview_files = self._file_explorer.get_selected_files()
        if not preview_files:
            messagebox.showwarning(
                "Uyarı",
                "Önce gezginden dosya seçin veya 'Tüm Dosyalara Uygula' seçeneğini açın."
            )
            self._opts["excel_rename"].set(False)
            return

        modal = ExcelRenameModal(self, preview_files)
        self.wait_window(modal)
        if not modal.result:
            self._opts["excel_rename"].set(False)
            self._excel_rename_params = None
            return

        self._excel_rename_params = {
            "excel_path": modal.excel_path,
            "sheet_name": modal.sheet_name,
            "column_letter": modal.column_letter,
            "start_num": modal.start_num,
            "output_column": modal.output_column,
            "mode": modal.mode,
        }

        # Modaldeki düğme "Onayla ve Çalıştır" dediği için kullanıcıdan
        # ana ekrandaki Uygula düğmesine ikinci kez basmasını beklemeyelim.
        # Checkbox callback'i tamamlandıktan sonra ana işlem akışını başlat.
        self.after(0, self._run)

    # ── Ana çalıştır ──────────────────────────────
    def _run(self):
        folder = self._folder.get().strip()
        if not folder or not os.path.isdir(folder):
            messagebox.showerror("Hata", "Lütfen önce geçerli bir klasör seçin.")
            return
            
        if self._apply_to_all.get():
            selected_files = [os.path.join(folder, f) for f in os.listdir(folder) if os.path.isfile(os.path.join(folder, f))]
            if not selected_files:
                messagebox.showwarning("Uyarı", "Seçili klasörde hiçbir dosya bulunamadı!")
                return
        else:
            selected_files = self._file_explorer.get_selected_files()
            if not selected_files:
                messagebox.showwarning("Uyarı", "Gezginden hiçbir dosya seçmediniz!\n\nİşlem yapılacak dosyaları listeden seçin veya 'Tüm Dosyalara Uygula' özelliğini açın.")
                return

        selected_operations = [k for k, v in self._opts.items() if v.get()]
        if not selected_operations:
            messagebox.showwarning("Uyarı", "Lütfen soldan en az bir işlem seçin.")
            return

        # ── İşlem Öncesi Kontroller & Önizlemeler ──
        if "truncate" in selected_operations:
            try:
                n = int(self._n_var.get().strip())
                if n <= 0: raise ValueError 
            except ValueError:
                messagebox.showerror("Hata", "Lütfen 0'dan büyük geçerli bir kesme sayısı (sayısal) giriniz.")
                return
                
            # Modal önizleme çağır (Önizleme boyunca ana ekran bekler)
            preview = TruncatePreviewModal(self, selected_files, n)
            self.wait_window(preview)
            
            if not preview.result:
                self._log("✖ Kesme işlemi kullanıcı tarafından iptal edildi.")
                return # Tüm run döngüsünü bitir

        if "seq_rename" in selected_operations:
            try:
                seq_start = int(self._seq_start_var.get().strip())
                if seq_start < 0: raise ValueError
            except ValueError:
                messagebox.showerror("Hata", "Lütfen geçerli bir başlangıç numarası (örneğin: 1) giriniz.")
                return

        # Excel ayarları checkbox açılırken toplanır.
        excel_params = self._excel_rename_params
        if "excel_rename" in selected_operations:
            if not excel_params:
                messagebox.showwarning(
                    "Uyarı",
                    "Excel ile numaralandırma ayarları eksik. Seçeneği kapatıp yeniden açın."
                )
                return
                
        # Run worker thread
        threading.Thread(target=self._run_worker, args=(selected_files, selected_operations, excel_params), daemon=True).start()

    def _run_worker(self, selected_files, selected_operations, excel_params=None):
        self._set_status("⏳ Çalışıyor...", WARNING)
        total = 0
        self._log("─" * 50)
        self._log(f"🚀 Başlatıldı → Seçilen {len(selected_files)} dosya işlenecek.")
        self._log("─" * 50)
        self._tab_view.set("İşlem Günlüğü")

        try:
            if "truncate" in selected_operations:
                n = int(self._n_var.get().strip()) # Validate yukarıda yapıldı
                self._log(f"\n[1] İlk {n} karakter sonrası siliniyor...")
                c = op_truncate_after_n(selected_files, n, self._log)
                self._log(f"   ✅ {c} dosya işlendi.")
                total += c

            if "jpeg2jpg" in selected_operations:
                self._log("\n[2] .jpeg → .jpg yeniden adlandırılıyor...")
                c = op_rename_ext(selected_files, "jpeg", "jpg", self._log)
                self._log(f"   ✅ {c} dosya işlendi.")
                total += c

            if "bmp2jpg" in selected_operations:
                self._log("\n[3] .bmp → .jpg dönüştürülüyor...")
                c = op_convert_to_jpg(selected_files, "bmp", self._log)
                self._log(f"   ✅ {c} dosya işlendi.")
                total += c

            if "png2jpg" in selected_operations:
                self._log("\n[4] .png → .jpg dönüştürülüyor...")
                c = op_convert_to_jpg(selected_files, "png", self._log)
                self._log(f"   ✅ {c} dosya işlendi.")
                total += c
                
            if "jfif2jpg" in selected_operations:
                self._log("\n[5] .jfif → .jpg dönüştürülüyor...")
                c = op_convert_to_jpg(selected_files, "jfif", self._log)
                self._log(f"   ✅ {c} dosya işlendi.")
                total += c

            if "pdf2jpg" in selected_operations:
                try:
                    dpi = int(self._dpi_var.get().strip())
                except ValueError:
                    self._log("\n✖ HATA: Lütfen geçerli bir DPI çözünürlüğü (sayısal) giriniz.")
                    self._set_status("Hata oluştu", DANGER)
                    return
                self._log(f"\n[6] .pdf → .jpg dönüştürülüyor (DPI={dpi})...")
                c = op_pdf_to_jpg(selected_files, dpi, self._log)
                self._log(f"   ✅ {c} sayfa işlendi.")
                total += c
                
            if "ext_photos" in selected_operations:
                self._log(f"\n[7] Yapay Zeka (Haar Cascade) ile vesikalıklar tespit edilip kesiliyor...")
                c = op_extract_photos(selected_files, self._log)
                self._log(f"   ✅ {c} adet yüz başarıyla ayrı dosyalara çıkarıldı.")
                total += c
            
            if "ext_photos_dlib" in selected_operations:
                self._log(f"\n[7b] Yapay Zeka (Dlib) ile vesikalıklar tespit edilip kesiliyor...")
                c = op_extract_photos_dlib(selected_files, self._log)
                self._log(f"   ✅ {c} adet yüz başarıyla ayrı dosyalara çıkarıldı.")
                total += c
            
            if "ext_photos_hybrid" in selected_operations:
                try:
                    hybrid_start = int(self._hybrid_start_var.get().strip())
                    if hybrid_start < 1: raise ValueError
                except ValueError:
                    self._log("\n✖ HATA: Hibrit başlangıç numarası geçersiz. Lütfen 1 veya daha büyük bir sayı girin.")
                    self._set_status("Hata oluştu", DANGER)
                    return
                self._log(f"\n[7c] Yapay Zeka (Hibrit: Haar + Dlib) ile vesikalıklar tespit edilip kesiliyor... (Başlangıç: {hybrid_start})")
                c = op_extract_photos_hybrid(selected_files, self._log, start_num=hybrid_start)
                self._log(f"   ✅ {c} adet yüz başarıyla ayrı dosyalara çıkarıldı.")
                total += c
                
            if "seq_rename" in selected_operations:
                seq_start = int(self._seq_start_var.get().strip())
                self._log(f"\n[8] Seçili dosyalar sırayla ({seq_start}'den başlayarak) baştan adlandırılıyor...")
                c = op_rename_sequence(selected_files, seq_start, self._log)
                self._log(f"   ✅ {c} dosya yeniden adlandırıldı.")
                total += c

            if "excel_rename" in selected_operations and excel_params:
                self._log(f"\n[9] 📊 Excel ile numaralandırma başlatılıyor...")
                c = op_excel_rename(
                    selected_files,
                    excel_params["excel_path"],
                    excel_params["sheet_name"],
                    excel_params["column_letter"],
                    excel_params["start_num"],
                    self._log,
                    excel_params["output_column"],
                    excel_params["mode"]
                )
                self._log(f"   ✅ {c} dosya Excel eşleştirmesiyle yeniden adlandırıldı.")
                total += c

        except Exception as e:
            self._log(f"\n✖ Beklenmeyen hata: {e}")
            self._set_status("Hata oluştu", DANGER)
            return

        self._log("─" * 50)
        self._log(f"🏁 Tamamlandı. Toplam {total} işlem yapıldı.")
        self._log("─" * 50 + "\n")
        self._set_status(f"✅ Tamamlandı — {total} işlem", SUCCESS)
        
        # After operation, refresh the explorer on the main thread in case files were renamed/deleted
        # also deselect operations
        self.after(500, self._post_run_cleanup)

    def _post_run_cleanup(self):
        self._refresh_file_list()
        self._deselect_all()
